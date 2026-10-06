from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

# Pinterest Aesthetic Palette
BG_SAND = RGBColor(244, 239, 234)
TERRACOTTA = RGBColor(200, 109, 81)
CHARCOAL = RGBColor(44, 44, 44)
SAGE = RGBColor(74, 107, 93)

slides_data = [
    {
        "title": "TripMate Kenya",
        "subtitle": "Revolutionizing Domestic & Gen Z Travel Through Tech",
        "bullets": [
            "Tagline: Discover Kenya. Connect. Travel Smarter.",
            "Founder & Lead Developer: Quinter Moraa Ondoro",
            "Location: Kisii / Nairobi, Kenya | Date: October 2026"
        ]
    },
    {
        "title": "Slide 2: The Problem",
        "subtitle": "Challenges Facing Young Domestic Travelers",
        "bullets": [
            "Fragmented Planning: Spending 5+ hours cross-referencing TikTok, blogs, and WhatsApp.",
            "Chaotic Group Logistics: Unpaid split bills, manual calculations, and payment delays.",
            "Overpriced Middlemen: Up to 35% markups from traditional agencies.",
            "Unseen Local Gems: Community staycations lack digital visibility."
        ]
    },
    {
        "title": "Slide 3: The Solution",
        "subtitle": "An All-In-One Social & Booking Experience",
        "bullets": [
            "Smart Curated Hub: Personalised itineraries tailored to budget and eco-preferences.",
            "Native Group Split-Pay: Integrated MPesa wallet for real-time expense sharing.",
            "Verified SME Marketplace: Direct listings with local guides with zero hidden fees.",
            "Social Travel Match: Connect with verified peer travelers heading to similar destinations."
        ]
    },
    {
        "title": "Slide 4: Target Market & Persona",
        "subtitle": "Primary Focus: Gen Z & Millennial Explorers",
        "bullets": [
            "Demographics: Tech-savvy youth aged 18–35 in Nairobi, Kisii, Mombasa, and Nakuru.",
            "Core Need: Transparent 'Under KES 5,000' weekend getaways.",
            "Key Motivation: Peer recommendations, authentic experiences, and social sharing."
        ]
    },
    {
        "title": "Slide 5: Market Opportunity & Kenya Tourism Stats",
        "subtitle": "Proven Local and Sector Demand",
        "bullets": [
            "5.15 Million Domestic Bed Nights recorded (Strong internal travel market).",
            "KES 452.2 Billion Inbound Earnings across 2.39M international arrivals.",
            "118% Mobile Penetration: 66M+ mobile connections driving app usage."
        ]
    },
    {
        "title": "Slide 6: Business & Revenue Model",
        "subtitle": "Diversified Monetization Streams",
        "bullets": [
            "Booking Commissions: 10%–12% take rate on stays and guided experiences.",
            "SME Subscriptions: KES 2,500/month flat fee for featured business placement.",
            "In-App Brand Partnerships: Sponsored posts for gear brands & local venues.",
            "TripMate Plus: KES 300/month premium subscription for offline maps & zero fees."
        ]
    },
    {
        "title": "Slide 7: Competition Comparison",
        "subtitle": "Positioned for Youth & Group Travel",
        "bullets": [
            "TripMate Kenya: Gen Z domestic focus, native MPesa group split pay, budget-friendly.",
            "Traditional Agencies: Safari focused, high agent fees, manual payments.",
            "Global Platforms (Airbnb/Booking): Hotel/stay focused, no group split tools."
        ]
    },
    {
        "title": "Slide 8: Marketing & User Acquisition",
        "subtitle": "4-Step Growth Strategy",
        "bullets": [
            "1. Micro-Influencer Campaigns: TikTok/Instagram #TembeaKenya challenges.",
            "2. Campus Ambassador Network: Student leaders at Kisii Univ, UoN, KU, and Strathmore.",
            "3. Community Challenges: 'Weekend Under KES 5k' user content contests.",
            "4. Mobility Partners: Strategic deals with SGR and regional shuttle operators."
        ]
    },
    {
        "title": "Slide 9: The Team",
        "subtitle": "Execution & Growth Mindset",
        "bullets": [
            "Quinter Moraa Ondoro — Founder & Lead Developer (Strategy & Engineering)",
            "Co-Founder & CTO (Mobile Architecture & Systems)",
            "Head of Marketing & Growth (Campus Networks & Community)",
            "Operations & Partner Lead (Vendor Onboarding & Finance)"
        ]
    },
    {
        "title": "Slide 10: Launch Budget & Year 1 Targets",
        "subtitle": "Financial Allocation & Growth Goals",
        "bullets": [
            "Budget Allocation (Total: KES 1,500,000):",
            " - 40% (KES 600,000) App Development & Infrastructure",
            " - 30% (KES 450,000) User Acquisition & Marketing",
            " - 20% (KES 300,000) SME Onboarding & Legal",
            " - 10% (KES 150,000) Contingency Reserve",
            "Year 1 Targets: 50,000+ Downloads | 250+ SME Vendors | KES 12M GBV"
        ]
    },
    {
        "title": "Slide 11: The Ask / Closing",
        "subtitle": "Join Us in Digitizing Kenya Domestic Travel",
        "bullets": [
            "Funding Ask: KES 1,500,000 Seed Capital / Grant.",
            "Use of Funds: App release, university campaigns, and onboarding 200+ SMEs across 10 counties.",
            "Contact: quinter@tripmate.co.ke | +254 700 000 000 | www.tripmate.co.ke"
        ]
    },
    {
        "title": "Slide 12: Data Sources & References",
        "subtitle": "Grounded in Sector Reports",
        "bullets": [
            "Tourism Research Institute (TRI) & Ministry of Tourism Performance Reports.",
            "Communications Authority of Kenya (CA) Mobile Statistics Report.",
            "DataReportal Kenya Digital Sector Overview."
        ]
    }
]

for item in slides_data:
    slide = prs.slides.add_slide(blank_layout)
    
    bg = slide.shapes.add_shape(1, Inches(0.5), Inches(0.5), Inches(12.333), Inches(6.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG_SAND
    bg.line.fill.background()
    
    tb_title = slide.shapes.add_textbox(Inches(1.0), Inches(0.8), Inches(11.333), Inches(1.2))
    tf_title = tb_title.text_frame
    tf_title.word_wrap = True
    
    p_title = tf_title.paragraphs[0]
    p_title.text = item["title"]
    p_title.font.size = Pt(28)
    p_title.font.bold = True
    p_title.font.color.rgb = TERRACOTTA
    
    p_sub = tf_title.add_paragraph()
    p_sub.text = item["subtitle"]
    p_sub.font.size = Pt(16)
    p_sub.font.color.rgb = SAGE
    
    tb_body = slide.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(11.333), Inches(4.3))
    tf_body = tb_body.text_frame
    tf_body.word_wrap = True
    
    for idx, b in enumerate(item["bullets"]):
        p = tf_body.paragraphs[0] if idx == 0 else tf_body.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(18)
        p.font.color.rgb = CHARCOAL
        p.space_after = Pt(14)

prs.save("TripMate_Kenya_Version2_PitchDeck.pptx")
print("Successfully generated TripMate_Kenya_Version2_PitchDeck.pptx!")